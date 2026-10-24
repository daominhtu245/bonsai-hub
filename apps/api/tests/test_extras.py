"""Tính năng bổ sung: lịch chăm sóc (Excel, .ics), mã quà tặng, lịch sử lượt, thông báo chung, tạm dừng,
gói bán chỉnh được, nhật ký ảnh cây, gửi thông báo hàng loạt, xuất Excel quản trị."""

import io
import uuid
from datetime import UTC, date, datetime, timedelta
from urllib.parse import urlparse

import pytest
from openpyxl import load_workbook

from bonsai_api.db.models import DesignJob, PhotoSource, PushSubscription, Stage, TreePhoto, User
from bonsai_api.services.care_calendar import add_months, months_from_text
from bonsai_api.settings import get_settings

from test_week5 import _admin, _done_job, bucket  # noqa: F401


@pytest.mark.parametrize("text,expected", [
    ("Khoảng 2 năm", 24), ("6–12 tháng", 6), ("Sau buổi tạo khung", 0), ("Hôm nay", 0), ("Khoảng 1,5 năm", 18),
    (None, 7), ("không rõ", 7),
])
def test_months_from_text(text, expected):
    assert months_from_text(text, 7) == expected


def test_add_months_clamps_day():
    assert add_months(date(2026, 1, 31), 1) == date(2026, 2, 28)
    assert add_months(date(2026, 11, 15), 14) == date(2028, 1, 15)


def _with_stages(db, job):
    for code, when, title in [("S1", "Sau buổi tạo khung", "Tạo khung"), ("S2", "Khoảng 1 năm", "Nuôi cành"),
                              ("S3", "Khoảng 3 năm", "Tạo tầng tán"), ("S4", "Khoảng 4 năm", "Hoàn thiện")]:
        db.add(Stage(job_id=job.id, code=code, time_from_start=when,
                     guide_vi={"title_vi": title, "actions_vi": [f"Việc {code}"], "care_vi": ["Tưới đều"], "best_time_vi": "Tháng 3–5"}))
    job.finished_at = datetime(2026, 10, 1, tzinfo=UTC)
    db.commit()


def test_care_plan_excel_and_calendar(db, client, login, bucket):  # noqa: F811
    headers, job = _done_job(db, client, login, bucket)
    _with_stages(db, job)
    links = client.post(f"/api/v1/jobs/{job.id}/exports", headers=headers).json()
    xurl, iurl = urlparse(links["xlsx_url"]), urlparse(links["ics_url"])
    r = client.get(f"{xurl.path}?{xurl.query}")  # không cần đăng nhập: link ký sẵn
    assert r.status_code == 200 and "lich-cham-soc" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    cells = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert "Việc: Việc S2" in cells and "01/10/2027" in cells and "☐" in cells  # S2 = +1 năm
    ics = client.get(f"{iurl.path}?{iurl.query}").text
    assert ics.count("BEGIN:VEVENT") == 4 + 3  # S1–S4 + 3 lần kiểm tra dây
    assert "DTSTART;VALUE=DATE:20291001" in ics and "TRIGGER:-P1D" in ics
    # chữ ký sai / hết hạn / người khác
    assert client.get(f"{xurl.path}?exp={xurl.query.split('&')[0][4:]}&sig=deadbeef").status_code == 403
    assert client.get(f"{xurl.path}?exp=1&sig=x").status_code == 403
    other, _ = login("khac@example.com")
    assert client.post(f"/api/v1/jobs/{job.id}/exports", headers=other).status_code == 404


def test_gift_codes_and_credit_history(db, client, login, monkeypatch):
    admin = _admin(client, login, monkeypatch)
    r = client.post("/api/v1/admin/gift-codes", headers=admin, json={"code": "tet 2027", "credits": 2, "max_uses": 1})
    assert r.status_code == 201 and r.json()["code"] == "TET2027" and r.json()["status"] == "active"
    auto = client.post("/api/v1/admin/gift-codes", headers=admin, json={"credits": 1}).json()
    assert auto["code"].startswith("BONSAI-")
    assert client.post("/api/v1/admin/gift-codes", headers=admin, json={"code": "TET2027", "credits": 1}).status_code == 409

    user, body = login("khach@example.com")
    r = client.post("/api/v1/me/gift-code", headers=user, json={"code": " tet2027 "})
    assert r.status_code == 200 and r.json() == {"credits_added": 2, "credits": 3}  # 1 lượt chào mừng + 2
    assert client.post("/api/v1/me/gift-code", headers=user, json={"code": "TET2027"}).json()["detail"] == "Bạn đã dùng mã này rồi."
    user2, _ = login("khach2@example.com")
    assert client.post("/api/v1/me/gift-code", headers=user2, json={"code": "TET2027"}).json()["detail"] == "Mã này đã hết lượt dùng."
    assert client.post("/api/v1/me/gift-code", headers=user2, json={"code": "KHONGCO"}).status_code == 400
    client.put(f"/api/v1/admin/gift-codes/{auto['code']}", headers=admin, json={"expires_at": (datetime.now(UTC) - timedelta(days=1)).isoformat()})
    assert client.post("/api/v1/me/gift-code", headers=user2, json={"code": auto["code"]}).json()["detail"] == "Mã này đã hết hạn."

    hist = client.get("/api/v1/me/credits", headers=user).json()
    assert [(h["label"], h["delta"]) for h in hist] == [("Mã quà tặng TET2027", 2), ("Quà chào mừng", 1)]
    codes = {g["code"]: g for g in client.get("/api/v1/admin/gift-codes", headers=admin).json()}
    assert codes["TET2027"]["status"] == "used_up" and codes[auto["code"]]["status"] == "expired"
    red = client.get("/api/v1/admin/gift-codes/TET2027/redemptions", headers=admin).json()
    assert [x["email"] for x in red] == ["khach@example.com"]


def test_ops_pause_announcement_and_trial(db, client, login, monkeypatch):
    admin = _admin(client, login, monkeypatch)
    assert client.get("/api/v1/announcement").json() is None
    r = client.put("/api/v1/admin/ops", headers=admin, json={
        "announcement": {"active": True, "text": "Tặng 1 lượt dịp Tết", "tone": "success"},
        "pause_analysis": True, "trial_credits": 3})
    assert r.status_code == 200 and r.json()["pause_analysis"] is True
    assert client.get("/api/v1/announcement").json() == {"text": "Tặng 1 lượt dịp Tết", "tone": "success"}
    user, body = login("moi@example.com")
    assert body["user"]["credits"] == 3  # lượt tặng khi đăng ký đổi được lúc chạy
    tree = client.post("/api/v1/trees", headers=user, json={"photos": [{"key": f"uploads/{body['user']['id']}/a.jpg", "view_index": 0, "is_front": True}]}).json()
    r = client.post(f"/api/v1/trees/{tree['id']}/jobs", headers=user, json={})
    assert r.status_code == 503 and "bảo trì" in r.json()["detail"]
    assert client.put("/api/v1/admin/ops", headers=admin, json={"usd_vnd_rate": 5}).status_code == 422
    audit = client.get("/api/v1/admin/audit?action=ops", headers=admin).json()
    assert audit["total"] == 1 and set(audit["items"][0]["detail"]) == {"announcement", "pause_analysis", "trial_credits"}


def test_packages_editable(db, client, login, monkeypatch):
    admin = _admin(client, login, monkeypatch)
    r = client.put("/api/v1/admin/packages/player", headers=admin, json={"amount_vnd": 179_000, "credits": 6})
    assert r.status_code == 200 and next(p for p in r.json() if p["code"] == "player")["amount_vnd"] == 179_000
    client.put("/api/v1/admin/packages/le", headers=admin, json={"active": False})
    pub = {p["code"]: p for p in client.get("/api/v1/packages").json()}
    assert "le" not in pub and pub["player"]["credits"] == 6
    client.put("/api/v1/admin/packages/nursery", headers=admin, json={"active": False})
    assert client.put("/api/v1/admin/packages/player", headers=admin, json={"active": False}).status_code == 400  # còn ≥ 1 gói
    monkeypatch.setattr(get_settings(), "bank_bin", "970422")
    monkeypatch.setattr(get_settings(), "bank_account", "0123456789")
    user, _ = login("mua@example.com")
    assert client.post("/api/v1/payments", headers=user, json={"package_code": "le"}).status_code == 400  # gói đã tắt
    p = client.post("/api/v1/payments", headers=user, json={"package_code": "player"}).json()
    assert p["amount_vnd"] == 179_000
    from bonsai_api.db.models import Payment

    assert db.get(Payment, uuid.UUID(p["id"])).credits == 6  # chốt số lượt lúc tạo đơn


def test_progress_photos_excluded_from_analysis(db, client, login, bucket, monkeypatch):  # noqa: F811
    sent = []
    monkeypatch.setattr("bonsai_api.queue.enqueue", lambda *a, **kw: sent.append(a) or True)
    headers, job = _done_job(db, client, login, bucket)
    uid = db.get(DesignJob, job.id).user_id
    r = client.post(f"/api/v1/trees/{job.tree_id}/progress", headers=headers, json={"key": f"uploads/{uid}/p1.jpg"})
    assert r.status_code == 201 and sent[-1][0] == "bonsai_worker.tasks.preprocess.normalize_tree_photos"
    assert client.post(f"/api/v1/trees/{job.tree_id}/progress", headers=headers, json={"key": "uploads/khac/x.jpg"}).status_code == 400
    lst = client.get(f"/api/v1/trees/{job.tree_id}/progress", headers=headers).json()
    assert len(lst) == 1 and lst[0]["ready"] is False
    photo = db.get(TreePhoto, uuid.UUID(lst[0]["id"]))
    assert photo.source == PhotoSource.progress
    # ảnh nhật ký không thành ảnh phôi của phiên
    j = client.get(f"/api/v1/jobs/{job.id}", headers=headers).json()
    assert all(p["id"] != str(photo.id) for p in j["photos"])
    assert client.delete(f"/api/v1/trees/{job.tree_id}/progress/{photo.id}", headers=headers).status_code == 204
    assert client.get(f"/api/v1/trees/{job.tree_id}/progress", headers=headers).json() == []


def test_broadcast_segments(db, client, login, monkeypatch):
    sent = []
    monkeypatch.setattr("bonsai_api.routers.admin_ops.queue.enqueue", lambda *a, **kw: sent.append(a) or True)
    admin = _admin(client, login, monkeypatch)
    _, body = login("nghe@example.com")
    payload = {"segment": "all", "title": "Mùa cắt tỉa", "body": "Đầu mùa mưa là lúc tốt để cắt tỉa si, sanh.", "path": "/vuon"}
    assert client.post("/api/v1/admin/broadcast", headers=admin, json=payload).status_code == 400  # chưa ai bật
    db.add(PushSubscription(user_id=uuid.UUID(body["user"]["id"]), endpoint="https://push/x", p256dh="k", auth="a"))
    db.commit()
    seg = {s["key"]: s["recipients"] for s in client.get("/api/v1/admin/broadcast/segments", headers=admin).json()}
    assert seg["all"] == 1 and seg["nursery"] == 0 and seg["never_paid"] == 1
    r = client.post("/api/v1/admin/broadcast", headers=admin, json=payload)
    assert r.status_code == 202 and r.json() == {"recipients": 1}
    assert sent[-1] == ("bonsai_worker.tasks.notify.broadcast", "all", "Mùa cắt tỉa", payload["body"], "/vuon")


def test_admin_export_xlsx(db, client, login, bucket, monkeypatch):  # noqa: F811
    _done_job(db, client, login, bucket)
    admin = _admin(client, login, monkeypatch)
    for kind, header in (("jobs", "Mã phiên"), ("users", "Email"), ("payments", "Nội dung CK")):
        r = client.get(f"/api/v1/admin/export/{kind}.xlsx", headers=admin)
        assert r.status_code == 200 and r.headers["content-disposition"].startswith('attachment; filename="bonsai-')
        ws = load_workbook(io.BytesIO(r.content)).active
        assert header in [c.value for c in ws[1]]
    jobs = load_workbook(io.BytesIO(client.get("/api/v1/admin/export/jobs.xlsx?status=done", headers=admin).content)).active
    assert jobs.max_row == 2  # tiêu đề + 1 phiên xong
    users_ws = load_workbook(io.BytesIO(client.get("/api/v1/admin/export/users.xlsx", headers=admin).content)).active
    assert users_ws.max_row >= 3
    assert db.query(User).count() >= 2
