"""Lịch chăm sóc cho người dùng: ngày dự kiến từng giai đoạn, xuất Excel (.xlsx) và lịch điện thoại (.ics).

Mốc thời gian lấy từ lộ trình (`stages.time_from_start`, vd. "Khoảng 2 năm", "6–12 tháng"), tính từ ngày có kết quả.
Chỉ là gợi ý: tốc độ thật phụ thuộc cây và cách chăm.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta

from ..schemas import ResultFull

DEFAULT_MONTHS = {"S0": 0, "S1": 0, "S2": 12, "S3": 30}
WIRE_CHECK_WEEKS = (3, 6, 10)  # nhắc kiểm tra dây uốn sau buổi tạo khung


def months_from_text(text: str | None, fallback: int) -> int:
    """"Khoảng 2 năm" → 24; "6–12 tháng" → 6; "Sau buổi tạo khung" / "Hôm nay" → 0."""
    if not text:
        return fallback
    t = text.lower()
    if "hôm nay" in t or "tạo khung" in t or "ngay" in t:
        return 0
    m = re.search(r"(\d+(?:[.,]\d+)?)", t)
    if not m:
        return fallback
    n = float(m.group(1).replace(",", "."))
    return round(n * 12) if "năm" in t else round(n) if "tháng" in t else fallback


def add_months(d: date, months: int) -> date:
    y, m = divmod(d.month - 1 + months, 12)
    year, month = d.year + y, m + 1
    for day in (d.day, 28):  # 31/1 + 1 tháng → 28/2
        try:
            return date(year, month, day)
        except ValueError:
            continue
    return date(year, month, 28)


@dataclass
class StageDate:
    code: str
    name: str
    when: str
    due: date
    actions: list[str]
    care: list[str]
    ready: list[str]
    mistakes: list[str]
    best_time: str


def stage_dates(result: ResultFull, start: date) -> list[StageDate]:
    years = result.years or 5
    out = []
    for s in result.stages:
        g = s.guide
        fallback = years * 12 if s.code == "S4" else DEFAULT_MONTHS.get(s.code, 0)
        months = months_from_text(s.time_from_start, fallback)
        out.append(StageDate(
            code=s.code, name=(g.title_vi if g and g.title_vi else s.name_vi), when=s.time_from_start or "",
            due=add_months(start, months), actions=g.actions_vi if g else [], care=g.care_vi if g else [],
            ready=g.ready_signs_vi if g else [], mistakes=g.common_mistakes_vi if g else [],
            best_time=g.best_time_vi if g else "",
        ))
    return out


def result_start(finished_at: datetime | None) -> date:
    return (finished_at or datetime.now(UTC)).astimezone(UTC).date()


# ---------- Excel ----------

def care_plan_xlsx(result: ResultFull, start: date, *, species_vi: str | None = None) -> bytes:
    """Bảng chữ to, dễ in: mỗi việc một dòng, có cột "Đã làm" để đánh dấu tay."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Lịch chăm sóc"
    thin = Side(style="thin", color="D4D2CC")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    head_fill = PatternFill("solid", fgColor="E8F0EA")
    stage_fill = PatternFill("solid", fgColor="F1F1EE")
    wrap = Alignment(wrap_text=True, vertical="top")

    name = result.tree_nickname or "Cây của tôi"
    ws["A1"] = f"Lịch chăm sóc — {name}"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = " · ".join(x for x in [species_vi, f"Thế {result.style_vi}" if result.style_vi else None,
                                        f"Bắt đầu {start.strftime('%d/%m/%Y')}"] if x)
    ws["A2"].font = Font(size=12, color="62615C")
    ws.append([])
    ws.append(["Giai đoạn", "Thời điểm", "Ngày dự kiến", "Việc cần làm", "Đã làm"])
    for c in ws[4]:
        c.font, c.fill, c.border = Font(size=12, bold=True), head_fill, border

    for sd in stage_dates(result, start):
        if sd.code == "S0" and not sd.actions:
            continue
        rows = [(f"Việc: {a}") for a in sd.actions] + [f"Chăm: {c}" for c in sd.care] + \
               [f"Sang giai đoạn sau khi: {r}" for r in sd.ready] + [f"Tránh: {m}" for m in sd.mistakes]
        if sd.best_time:
            rows.insert(0, f"Thời điểm tốt: {sd.best_time}")
        first = True
        for text in rows or ["—"]:
            ws.append([sd.name if first else "", sd.when if first else "", sd.due.strftime("%d/%m/%Y") if first else "", text, "☐"])
            for c in ws[ws.max_row]:
                c.font, c.alignment, c.border = Font(size=12, bold=first and c.column <= 3), wrap, border
                if first and c.column <= 3:
                    c.fill = stage_fill
            ws.cell(ws.max_row, 5).alignment = Alignment(horizontal="center", vertical="top")
            first = False

    ws.append([])
    ws.append([result.disclaimer_vi])
    ws.cell(ws.max_row, 1).font = Font(size=10, italic=True, color="62615C")
    for col, width in zip("ABCDE", (22, 18, 14, 70, 9)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- lịch điện thoại (.ics) ----------

def _ics_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def _fold(line: str) -> str:
    """RFC 5545: dòng ≤ 75 octet, xuống dòng + khoảng trắng."""
    raw = line.encode("utf-8")
    if len(raw) <= 75:
        return line
    parts, cur = [], b""
    for ch in line:
        b = ch.encode("utf-8")
        if len(cur) + len(b) > 74:
            parts.append(cur.decode("utf-8"))
            cur = b
        else:
            cur += b
    parts.append(cur.decode("utf-8"))
    return "\r\n ".join(parts)


def care_calendar_ics(result: ResultFull, start: date, *, url: str) -> bytes:
    """Mỗi giai đoạn một sự kiện cả ngày (nhắc trước 1 ngày) + nhắc kiểm tra dây uốn sau buổi tạo khung."""
    name = result.tree_nickname or "cây"
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    events: list[tuple[str, date, str, str]] = []
    for sd in stage_dates(result, start):
        if sd.code == "S0":
            continue
        desc = "\n".join([f"• {a}" for a in sd.actions] + ([f"Thời điểm tốt: {sd.best_time}"] if sd.best_time else []) + [url])
        events.append((f"{sd.code}", sd.due, f"Bonsai: {sd.name} — {name}", desc))
    s1 = next((sd.due for sd in stage_dates(result, start) if sd.code == "S1"), start)
    for w in WIRE_CHECK_WEEKS:
        events.append((f"wire{w}", s1 + timedelta(weeks=w), f"Bonsai: kiểm tra dây uốn — {name}",
                       f"Dây hằn vào vỏ thì tháo ngay, uốn lại nếu cành chưa giữ dáng.\n{url}"))
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Bonsai AI//Lich cham soc//VI", "CALSCALE:GREGORIAN",
             f"X-WR-CALNAME:{_ics_escape('Chăm ' + name)}"]
    for key, day, summary, desc in events:
        lines += [
            "BEGIN:VEVENT", f"UID:{result.job_id}-{key}@bonsai-ai", f"DTSTAMP:{stamp}",
            f"DTSTART;VALUE=DATE:{day.strftime('%Y%m%d')}", f"DTEND;VALUE=DATE:{(day + timedelta(days=1)).strftime('%Y%m%d')}",
            f"SUMMARY:{_ics_escape(summary)}", f"DESCRIPTION:{_ics_escape(desc)}", f"URL:{url}",
            "BEGIN:VALARM", "ACTION:DISPLAY", f"DESCRIPTION:{_ics_escape(summary)}", "TRIGGER:-P1D", "END:VALARM",
            "END:VEVENT",
        ]
    lines.append("END:VCALENDAR")
    return ("\r\n".join(_fold(x) for x in lines) + "\r\n").encode("utf-8")
