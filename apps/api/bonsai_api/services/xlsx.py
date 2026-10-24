"""Xuất bảng ra Excel (.xlsx) cho trang quản trị: tiêu đề in đậm, lọc, cố định dòng đầu, độ rộng cột vừa nội dung."""

import io
from collections.abc import Iterable, Sequence
from datetime import datetime


def write_xlsx(title: str, headers: Sequence[str], rows: Iterable[Sequence[object]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(list(headers))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8F0EA")
    widths = [len(str(h)) for h in headers]
    for r in rows:
        values = [v.replace(tzinfo=None) if isinstance(v, datetime) else v for v in r]  # Excel không nhận múi giờ
        ws.append(values)
        for i, v in enumerate(values):
            widths[i] = max(widths[i], min(60, len(str(v)) if v is not None else 0))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w + 2
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, datetime):
                c.number_format = "dd/mm/yyyy hh:mm"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
